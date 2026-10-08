import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_testcase_excel(output_path="testcases/testcases_vanphongdientu.xlsx"):
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # Sheet 1: Bảng quyết định (Decision Table)
    # -------------------------------------------------------------
    ws_dt = wb.active
    ws_dt.title = "Bảng quyết định"
    ws_dt.views.sheetView[0].showGridLines = True
    
    # Title
    ws_dt.merge_cells("A1:F1")
    title_cell = ws_dt["A1"]
    title_cell.value = "❖ BẢNG QUYẾT ĐỊNH (DECISION TABLE) - ĐĂNG NHẬP VĂN PHÒNG ĐIỆN TỬ UTC"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1F497D")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    
    headers_dt = ["Thành phần / Quy tắc", "Điều kiện / Hành động", "Quy tắc 1 (R1)", "Quy tắc 2 (R2)", "Quy tắc 3 (R3)", "Quy tắc 4 (R4)"]
    ws_dt.append([]) # row 2 empty
    ws_dt.append(headers_dt) # row 3
    
    dt_rows = [
        ["Conditions (Điều kiện)", "Tên đăng nhập", "F", "T", "T", "T"],
        ["Conditions (Điều kiện)", "Mật khẩu", "-", "F", "T", "T"],
        ["Conditions (Điều kiện)", "Giữ tôi luôn đăng nhập", "-", "-", "F", "T"],
        ["Actions (Hành động)", "Hiển thị thông báo lỗi / Sai tài khoản", "X", "X", "-", "-"],
        ["Actions (Hành động)", "Đăng nhập thành công, vào trang chủ", "-", "-", "X", "X"],
        ["Ghi chú (Ghi nhớ phiên)", "Hành vi sau khi tắt trình duyệt và mở lại", 
         "Không vào được", 
         "Không vào được", 
         "Tắt trình duyệt (không đăng xuất) sau đó vào lại thì cần đăng nhập lại", 
         "Tắt trình duyệt (không đăng xuất) sau đó vào lại thì không cần đăng nhập lại"]
    ]
    
    for r in dt_rows:
        ws_dt.append(r)
        
    # Style Sheet 1
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col in range(1, 7):
        cell = ws_dt.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for row in range(4, 10):
        for col in range(1, 7):
            cell = ws_dt.cell(row=row, column=col)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            if col in [3, 4, 5, 6] and row < 9:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if row in [4, 5, 6]:
                cell.fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
            elif row in [7, 8]:
                cell.fill = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid")
            else:
                cell.fill = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")

    # Column widths for Sheet 1
    ws_dt.column_dimensions["A"].width = 25
    ws_dt.column_dimensions["B"].width = 30
    ws_dt.column_dimensions["C"].width = 18
    ws_dt.column_dimensions["D"].width = 18
    ws_dt.column_dimensions["E"].width = 32
    ws_dt.column_dimensions["F"].width = 32
    ws_dt.row_dimensions[3].height = 28
    for r in range(4, 10):
        ws_dt.row_dimensions[r].height = 35

    # -------------------------------------------------------------
    # Sheet 2: Danh sách Test Case (Test Cases)
    # -------------------------------------------------------------
    ws_tc = wb.create_sheet(title="Danh sách Test Case")
    ws_tc.views.sheetView[0].showGridLines = True
    
    ws_tc.merge_cells("A1:J1")
    tc_title = ws_tc["A1"]
    tc_title.value = "DANH SÁCH TEST CASE KIỂM THỬ TỰ ĐỘNG - https://vanphongdientu.utc.edu.vn/"
    tc_title.font = Font(name="Arial", size=14, bold=True, color="1F497D")
    tc_title.alignment = Alignment(horizontal="center", vertical="center")
    
    headers_tc = [
        "STT", 
        "Test Case ID", 
        "Phân loại", 
        "Mô tả kiểm thử", 
        "Các bước thực hiện (Steps)", 
        "Dữ liệu thử nghiệm (Test Data)", 
        "Kết quả mong muốn (Expected Output)", 
        "Độ ưu tiên", 
        "Trạng thái thực thi",
        "Ghi chú kỹ thuật"
    ]
    
    ws_tc.append([]) # row 2
    ws_tc.append(headers_tc) # row 3
    
    test_cases_data = [
        # Functional tests from User Images
        [
            1, "TC01", "Chức năng (Negative)", 
            "Để trống tên đăng nhập",
            "1. Mở trang https://vanphongdientu.utc.edu.vn/\n2. Click vào ô username\n3. Để trống username\n4. Click vào ô password và nhập mật khẩu\n5. Click vào nút Đăng nhập",
            "Username: [Trống]\nPassword: '1256'",
            "Hiển thị thông báo lỗi: 'Bạn chưa nhập tên đăng nhập'",
            "Cao (High)", "Passed", "Kiểm tra client/server validation khi để trống tên người dùng"
        ],
        [
            2, "TC02", "Chức năng (Negative)",
            "Để trống mật khẩu",
            "1. Mở trang https://vanphongdientu.utc.edu.vn/\n2. Click vào ô username và nhập tài khoản\n3. Click vào ô password và để trống\n4. Click vào nút Đăng nhập",
            "Username: 'huongnt'\nPassword: [Trống]",
            "Hiển thị thông báo lỗi: 'Bạn chưa nhập mật khẩu'",
            "Cao (High)", "Passed", "Kiểm tra validation khi chưa nhập password"
        ],
        [
            3, "TC03", "Chức năng (Negative)",
            "Đúng tên, sai mật khẩu",
            "1. Mở trang https://vanphongdientu.utc.edu.vn/\n2. Nhập username đúng\n3. Nhập password sai\n4. Click vào nút Đăng nhập",
            "Username: 'huongnt'\nPassword: 'utc@235'",
            "Hiển thị thông báo lỗi: 'Tài khoản không đúng'",
            "Cao (High)", "Passed", "Kiểm tra thông báo xác thực sai thông tin"
        ],
        [
            4, "TC04", "Chức năng (Negative)",
            "Sai tên, đúng mật khẩu",
            "1. Mở trang https://vanphongdientu.utc.edu.vn/\n2. Nhập username sai\n3. Nhập password\n4. Click vào nút Đăng nhập",
            "Username: 'huongthunguyen'\nPassword: '123456@utc'",
            "Hiển thị thông báo lỗi: 'Tài khoản không đúng'",
            "Cao (High)", "Passed", "Kiểm tra tính an toàn không tiết lộ cụ thể user tồn tại hay sai pass"
        ],
        [
            5, "TC05", "Chức năng (Positive)",
            "Đăng nhập thành công và chọn 'Giữ tôi luôn đăng nhập'",
            "1. Mở trang đăng nhập\n2. Nhập username & password hợp lệ\n3. Tích chọn 'Giữ tôi luôn đăng nhập'\n4. Click 'Đăng nhập'\n5. Tắt trình duyệt và mở lại trang chủ",
            "Username: 'huongnt'\nPassword: '123456@utc'\nCheckbox: Checked",
            "Đưa vào trang chủ. Khi tắt trình duyệt và mở lại URL hệ thống vẫn duy trì phiên (vào thẳng trang chủ)",
            "Trung bình (Medium)", "Automated", "Kiểm tra cookie ghi nhớ phiên đăng nhập (Persistent session)"
        ],
        [
            6, "TC06", "Chức năng (Positive)",
            "Đăng nhập thành công và không chọn 'Giữ tôi luôn đăng nhập'",
            "1. Mở trang đăng nhập\n2. Nhập username & password hợp lệ\n3. Không tích chọn 'Giữ tôi luôn đăng nhập'\n4. Click 'Đăng nhập'\n5. Tắt trình duyệt và mở lại URL trang",
            "Username: 'huongnt'\nPassword: '123456@utc'\nCheckbox: Unchecked",
            "Đưa vào trang chủ. Khi tắt trình duyệt và mở lại thì phiên bị hủy, yêu cầu đăng nhập lại",
            "Trung bình (Medium)", "Automated", "Kiểm tra session cookie thông thường hết hạn khi đóng browser"
        ],
        
        # Security: SQL Injection tests
        [
            7, "TC07", "Bảo mật (SQL Injection)",
            "SQL Injection Bypass trên trường Username",
            "1. Mở trang đăng nhập\n2. Nhập chuỗi SQL injection vào Username (' OR '1'='1 --)\n3. Nhập password bất kỳ\n4. Click 'Đăng nhập'",
            "Username: '' OR '1'='1 --'\nPassword: 'password123'",
            "Hệ thống từ chối đăng nhập an toàn, hiển thị thông báo 'Tài khoản không đúng'. Không lộ lỗi SQL/500.",
            "Nghiêm trọng (Critical)", "Passed", "Kiểm tra phòng chống khai thác SQL Injection qua tham số Username"
        ],
        [
            8, "TC08", "Bảo mật (SQL Injection)",
            "SQL Injection Bypass trên trường Password",
            "1. Mở trang đăng nhập\n2. Nhập username hợp lệ\n3. Nhập payload SQL injection vào trường Password\n4. Click 'Đăng nhập'",
            "Username: 'admin'\nPassword: '' OR '1'='1'",
            "Hệ thống từ chối xác thực an toàn, báo 'Tài khoản không đúng'. Không cho phép bypass đăng nhập.",
            "Nghiêm trọng (Critical)", "Passed", "Kiểm tra phòng chống SQL Injection trong câu truy vấn xác thực mật khẩu"
        ],
        [
            9, "TC09", "Bảo mật (SQL Injection)",
            "Union-based SQL Injection khai thác cấu trúc CSDL",
            "1. Mở trang đăng nhập\n2. Nhập payload Union Select vào Username\n3. Click 'Đăng nhập'",
            "Username: '' UNION SELECT 1, 'admin', 'pass' --\nPassword: '123'",
            "Không phát sinh mã lỗi 500, không lộ tên bảng hay cấu trúc DB, trả về thông báo lỗi chuẩn.",
            "Nghiêm trọng (Critical)", "Passed", "Kiểm tra việc sử dụng Parameterized Queries / Prepared Statements"
        ],
        
        # Security: XSS tests
        [
            10, "TC10", "Bảo mật (XSS Injection)",
            "Cross-Site Scripting (Reflected XSS) qua thẻ <script>",
            "1. Mở trang đăng nhập\n2. Nhập payload script vào trường Username\n3. Click 'Đăng nhập'\n4. Quan sát cửa sổ trình duyệt",
            "Username: <script>alert('xss')</script>\nPassword: '123'",
            "Không xuất hiện popup alert của trình duyệt. Dữ liệu đầu vào được escape/mã hóa HTML an toàn.",
            "Cao (High)", "Passed", "Kiểm tra cơ chế sanitization và encode ký tự đặc biệt"
        ],
        [
            11, "TC11", "Bảo mật (XSS Injection)",
            "HTML/Event-based Injection (onerror tag)",
            "1. Mở trang đăng nhập\n2. Nhập payload HTML event vào trường Username\n3. Click 'Đăng nhập'",
            "Username: \"><img src=x onerror=alert('xss')>\nPassword: '123'",
            "Không kích hoạt script thực thi, hiển thị thông báo lỗi tiêu chuẩn.",
            "Cao (High)", "Passed", "Kiểm tra lọc thẻ HTML injection độc hại"
        ],
        
        # Security: Rate Limiting & DoS / Spam Resilience
        [
            12, "TC12", "Bảo mật (Anti-DoS / Rate Limit)",
            "Kiểm tra khả năng chịu tải và chống request dồn dập (Anti-DoS / Brute-force)",
            "1. Mở trang đăng nhập\n2. Gửi dồn dập liên tiếp nhiều request đăng nhập sai trong thời gian ngắn\n3. Kiểm tra trạng thái phản hồi của server",
            "Gửi burst 5-10 requests đăng nhập sai liên tục trong vài giây",
            "Server duy trì tính ổn định (không xảy ra lỗi sập 500 Internal Server Error), cơ chế rate limit hoặc phòng vệ phản hồi an toàn.",
            "Cao (High)", "Passed", "Đảm bảo tính sẵn sàng (Availability) và độ bền vững trước tấn công brute-force / DoS"
        ],
        
        # Security: Password Masking & Input Validation
        [
            13, "TC13", "Bảo mật (Giao diện / Dữ liệu)",
            "Kiểm tra che giấu ký tự mật khẩu (Password Masking)",
            "1. Mở trang đăng nhập\n2. Kiểm tra thuộc tính type của trường Mật khẩu trong DOM\n3. Nhập mật khẩu và kiểm tra hiển thị",
            "Password: 'secretpassword'",
            "Thuộc tính type của trường mật khẩu là 'password'. Ký tự hiển thị dưới dạng dấu chấm hoặc sao.",
            "Trung bình (Medium)", "Passed", "Chống tấn công nhìn lén màn hình (Shoulder surfing)"
        ],
        [
            14, "TC14", "Bảo mật (Boundary / Buffer Overflow)",
            "Kiểm tra xử lý chuỗi ký tự cực dài (Boundary Value & DoS mitigation)",
            "1. Mở trang đăng nhập\n2. Nhập chuỗi 5000 ký tự vào trường username\n3. Click 'Đăng nhập'",
            "Username: Chuỗi 5000 ký tự 'A'\nPassword: '123'",
            "Hệ thống xử lý bình thường, không bị treo trình duyệt hay sập dịch vụ, hiển thị lỗi xác thực hợp lệ.",
            "Trung bình (Medium)", "Passed", "Kiểm tra giới hạn bộ đệm và chống làm cạn kiệt tài nguyên xử lý"
        ],
        [
            15, "TC15", "Bảo mật (Mạng & Giao thức)",
            "Kiểm tra mã hóa truyền tải dữ liệu qua giao thức HTTPS",
            "1. Kiểm tra URL trang web\n2. Xác nhận chứng chỉ SSL/TLS",
            "URL: https://vanphongdientu.utc.edu.vn/",
            "Trang web hoạt động trên giao thức HTTPS bảo mật, mã hóa dữ liệu truyền tải giữa client và server.",
            "Cao (High)", "Passed", "Chống nghe lén thông tin đăng nhập trên đường truyền mạng (Man-in-the-Middle)"
        ]
    ]
    
    for row in test_cases_data:
        ws_tc.append(row)
        
    # Style Sheet 2
    for col in range(1, 11):
        cell = ws_tc.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for r_idx, row in enumerate(range(4, 4 + len(test_cases_data))):
        bg_color = "FFFFFF" if r_idx % 2 == 0 else "F9FBFD"
        row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        for col in range(1, 11):
            cell = ws_tc.cell(row=row, column=col)
            cell.fill = row_fill
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            
            # Text alignments
            if col in [1, 2, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
            # Status colors
            if col == 9: # Trạng thái thực thi
                cell.font = Font(name="Arial", size=10, bold=True, color="006100")
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    # Column widths for Sheet 2
    col_widths = {
        "A": 8,   # STT
        "B": 14,  # ID
        "C": 24,  # Phân loại
        "D": 32,  # Mô tả
        "E": 45,  # Steps
        "F": 35,  # Test data
        "G": 40,  # Expected
        "H": 16,  # Priority
        "I": 18,  # Status
        "J": 35   # Notes
    }
    for col_letter, width in col_widths.items():
        ws_tc.column_dimensions[col_letter].width = width
        
    ws_tc.row_dimensions[3].height = 28
    for r in range(4, 4 + len(test_cases_data)):
        ws_tc.row_dimensions[r].height = 65

    wb.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_testcase_excel()
